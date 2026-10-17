#!/usr/bin/env python3
"""Build tracker.xlsx: 8 tabs, header rows, Priority/Status dropdowns, Dashboard COUNTIF formulas. No data rows.

Run (openpyxl lives in the repo-local .venv, nothing global):
  uv venv .venv && uv pip install --python .venv/bin/python openpyxl
  .venv/bin/python tracker/build-xlsx.py [out.xlsx]        (default: tracker.xlsx in the current directory)
Google Sheets imports the file (File > Import).

APPROXIMATE: the tab list and column letters come from the sync ranges, not from a read of a live Sheet. Widths,
colours, the Inbox tab's column names (B-D, F, H, I) and the Agents M-N headers are a best guess; tune to taste.
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

import tabs

HEAD = PatternFill("solid", fgColor="D9E2F3")
MAXROW = 1000


def sheet_with_header(wb, name, cols):
    ws = wb.create_sheet(name)
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEAD
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s1" % tabs.col(len(cols))
    for i, h in enumerate(cols, 1):
        ws.column_dimensions[tabs.col(i)].width = max(12, min(48, len(h) + 6))
    return ws


def dropdown(ws, letter, values):
    dv = DataValidation(type="list", formula1='"%s"' % ",".join(values), allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("%s2:%s%d" % (letter, letter, MAXROW))


def build(out):
    wb = Workbook()
    dash = wb.active
    dash.title = "Dashboard"
    cols = {}
    for tab, csv_name, extra in tabs.TABS:
        cols[tab] = tabs.header(csv_name) + extra
        sheet_with_header(wb, tab, cols[tab])
    sheet_with_header(wb, "Inbox", tabs.INBOX_HEADERS)

    t = wb["Tasks"]
    pc = tabs.col(cols["Tasks"].index("Priority") + 1)
    sc = tabs.col(cols["Tasks"].index("Status") + 1)
    dropdown(t, pc, tabs.PRIORITIES)
    dropdown(t, sc, tabs.STATUSES)
    t.conditional_formatting.add("%s2:%s%d" % (pc, pc, MAXROW),
                                 CellIsRule(operator="equal", formula=['"P0"'], fill=PatternFill("solid", bgColor="F4CCCC")))
    t.conditional_formatting.add("%s2:%s%d" % (sc, sc, MAXROW),
                                 CellIsRule(operator="equal", formula=['"Done"'], fill=PatternFill("solid", bgColor="D9EAD3")))
    dc = tabs.col(cols["Decisions"].index("Status") + 1)
    dropdown(wb["Decisions"], dc, ["open", "done", "dropped"])

    dash["A1"] = "Dashboard"
    dash["A1"].font = Font(bold=True, size=14)
    dash["A3"], dash["B3"] = "Priority", "Tasks"
    dash["D3"], dash["E3"] = "Status", "Tasks"
    for c in ("A3", "B3", "D3", "E3"):
        dash[c].font = Font(bold=True)
        dash[c].fill = HEAD
    for i, p in enumerate(tabs.PRIORITIES):
        r = 4 + i
        dash.cell(r, 1, p)
        dash.cell(r, 2, "=COUNTIF(Tasks!$%s$2:$%s$%d,A%d)" % (pc, pc, MAXROW, r))
    for i, s in enumerate(tabs.STATUSES):
        r = 4 + i
        dash.cell(r, 4, s)
        dash.cell(r, 5, "=COUNTIF(Tasks!$%s$2:$%s$%d,D%d)" % (sc, sc, MAXROW, r))
    dash["A10"], dash["B10"] = "Open decisions", "=COUNTIF(Decisions!$%s$2:$%s$%d,\"open\")" % (dc, dc, MAXROW)
    dash["A11"], dash["B11"] = "Agents over 30% context", "=COUNTIF(Agents!$G$2:$G$%d,\"yes\")" % MAXROW
    dash["A12"], dash["B12"] = "Open agent briefs", "=COUNTIF('Agent inboxes'!$F$2:$F$%d,\"open\")" % MAXROW
    for c, w in (("A", 26), ("B", 10), ("D", 20), ("E", 10)):
        dash.column_dimensions[c].width = w
    wb.save(out)


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "tracker.xlsx"
    build(out)
    print("wrote", out)
